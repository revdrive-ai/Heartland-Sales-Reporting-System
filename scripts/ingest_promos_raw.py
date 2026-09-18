#!/usr/bin/env python3
"""Ingest the raw Telus retail-promotions export (single 'All Promos' sheet)
into the same fixtures ingest_promos.py builds from the transformed template.

Source: data/raw/Retail_Promotions_*.xlsx — the export Telus produces
directly, one row per component line, WITH the Dist Name column (the
distributor/ship-to that aligns lines to customer-crosswalk accounts and
splits the Safeway Mountain West planner into Safeway Denver / Safeway IMW).

Outputs (unchanged paths, so the app and loaders pick them up as-is):
  data/promos/promotions.json.gz   one row per promotion + dist_names[]
  data/promos/promo-lines.json.gz  one row per component line + dist_name
  data/promos/meta.json            snapshot, counts, reconciled totals

Differences from the transformed template handled here:
  - a trailing 'Sum:' footer row (no Planner ID) is dropped
  - Rate UOM arrives as Cases/Percentages/Eaches/blank → normalized to the
    Case/Percent/Each/Lump Sum vocabulary the app uses
  - channel / market / customer_code / planner_template don't exist in the
    raw export; they are functional per customer, so they carry forward from
    the previous fixture by customer_id
  - line_id keeps the promo|component|item recipe; lines that differ only by
    Dist Name (the Mountain West split) append |<dist> to stay unique

Run:  python3 scripts/ingest_promos_raw.py [path-to-xlsx]
"""
import gzip
import json
import pathlib
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parent.parent
RAW = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "raw" / "Retail_Promotions_08212026_2_1.xlsx"
OUT = ROOT / "data" / "promos"
SNAPSHOT_DATE = "2026-08-21"  # export of 08/21/2026 (filename), revision 2

UOM = {"Cases": "Case", "Percentages": "Percent", "Eaches": "Each", "": "Lump Sum"}


def money(v):
    if v is None or str(v).strip() in ("", "None"):
        return 0.0
    return round(float(str(v).replace("$", "").replace(",", "")), 2)


def iso(v):
    return str(v).strip()[:10]


def main() -> None:
    wb = load_workbook(RAW, read_only=True, data_only=True)
    it = wb["All Promos"].iter_rows(values_only=True)
    header = [str(h).strip() for h in next(it)]
    raw = [dict(zip(header, r)) for r in it if any(v is not None for v in r)]

    rows, dropped = [], 0
    for r in raw:
        if str(r.get("Planner ID") or "").strip() in ("", "None"):
            dropped += 1  # the 'Sum:' footer (and any stray blank)
            continue
        rows.append(r)
    print(f"rows: {len(rows)} (dropped {dropped} footer/blank)")

    errors = []

    # ---- promotions (headers are constant per promo — validated) -----------
    by_promo = defaultdict(list)
    for r in rows:
        by_promo[str(r["Promo ID Base"]).strip()].append(r)

    prev = {p["promo_id"]: p for p in json.load(gzip.open(OUT / "promotions.json.gz", "rt"))}
    by_cust = {}
    for p in prev.values():
        by_cust.setdefault(p["customer_id"], p)

    header_fields = ["Fiscal Year", "Template Type", "Promo Performance Type", "Promo Status",
                     "Promo Title", "Promo Perf Start", "Promo Perf End", "Planner ID", "Planner Title"]
    promotions, lines = [], []
    for pid, rs in sorted(by_promo.items()):
        for f in header_fields:
            if len({str(r[f]) for r in rs}) > 1:
                errors.append(f"{pid}: inconsistent {f!r} across lines")
        r0 = rs[0]
        cust_id = str(r0["Planner ID"]).strip()
        carried = prev.get(pid) or by_cust.get(cust_id) or {}
        if not carried:
            errors.append(f"{pid}: customer {cust_id} unknown — channel/market need a mapping")
        start, end = iso(r0["Promo Perf Start"]), iso(r0["Promo Perf End"])
        if end < start:
            errors.append(f"{pid}: end {end} before start {start}")
        dist_names = sorted({str(r["Dist Name"]).strip() for r in rs if str(r.get("Dist Name") or "").strip()})
        promotions.append({
            "promo_id": pid,
            "promo_title": str(r0["Promo Title"]).strip(),
            "fiscal_year": int(r0["Fiscal Year"]),
            "promo_status": str(r0["Promo Status"]).strip(),
            "template_type": str(r0["Template Type"]).strip(),
            "performance_type": str(r0["Promo Performance Type"]).strip(),
            "customer_id": cust_id,
            "customer_name": str(r0["Planner Title"]).strip(),
            "customer_code": carried.get("customer_code"),
            "channel": carried.get("channel"),
            "market": carried.get("market"),
            "planner_template": carried.get("planner_template", ""),
            "start_date": start,
            "end_date": end,
            "dist_names": dist_names,
            "line_count": len(rs),
            "planned_amount": round(sum(money(r["Promo $"]) for r in rs), 2),
            "actual_amount": round(sum(money(r["Actual Spend"]) for r in rs), 2),
        })

        # ---- lines — stable ids, |<dist> only where needed to disambiguate --
        key_counts = Counter(f'{pid}|{str(r["Component Name"]).strip()}|{str(r["Item Number"]).strip()}' for r in rs)
        for r in rs:
            base = f'{pid}|{str(r["Component Name"]).strip()}|{str(r["Item Number"]).strip()}'
            dist = str(r.get("Dist Name") or "").strip() or None
            uom_raw = str(r.get("Rate UOM") or "").strip()
            uom = UOM.get(uom_raw, uom_raw)
            if uom not in ("Case", "Each", "Percent", "Lump Sum"):
                errors.append(f"{pid}: unknown Rate UOM {uom_raw!r}")
            lines.append({
                "line_id": base + (f"|{dist}" if key_counts[base] > 1 else ""),
                "promo_id": pid,
                "component_type": str(r["Component Name"]).strip(),
                "brand": str(r["Brand"]).strip(),
                "item_number": str(r["Item Number"]).strip(),
                "item_description": (str(r["Item Desc"]).strip() or None) if r.get("Item Desc") else None,
                "rate": money(r["Component Rate"]),
                "rate_uom": uom,
                "planned_amount": money(r["Promo $"]),
                "actual_amount": money(r["Actual Spend"]),
                "dist_name": dist,
            })

    ids = Counter(l["line_id"] for l in lines)
    for k, n in ids.items():
        if n > 1:
            errors.append(f"duplicate line_id after dist suffix: {k} x{n}")

    if errors:
        print(f"\n✗ {len(errors)} validation errors:")
        for e in errors[:20]:
            print("  ·", e)
        sys.exit(1)

    # ---- write ---------------------------------------------------------------
    OUT.mkdir(parents=True, exist_ok=True)
    with gzip.open(OUT / "promotions.json.gz", "wt") as f:
        json.dump(promotions, f)
    with gzip.open(OUT / "promo-lines.json.gz", "wt") as f:
        json.dump(lines, f)
    meta = {
        "snapshot_date": SNAPSHOT_DATE,
        "source": RAW.name,
        "promotions": len(promotions),
        "lines": len(lines),
        "planned_total": round(sum(p["planned_amount"] for p in promotions), 2),
        "actual_total": round(sum(p["actual_amount"] for p in promotions), 2),
        "dist_names": sorted({d for p in promotions for d in p["dist_names"]}),
    }
    (OUT / "meta.json").write_text(json.dumps(meta, indent=2) + "\n")

    prev_ids = set(prev)
    now_ids = {p["promo_id"] for p in promotions}
    print(f"✓ {len(promotions)} promotions · {len(lines)} lines · planned ${meta['planned_total']:,.0f} · actual ${meta['actual_total']:,.0f}")
    print(f"  vs previous book: +{len(now_ids - prev_ids)} new promos · -{len(prev_ids - now_ids)} gone · {len(meta['dist_names'])} dist names")
    split = [p["promo_id"] for p in promotions if len(p["dist_names"]) > 1]
    print(f"  promos spanning multiple dists: {len(split)} (all {sorted({tuple(p['dist_names']) for p in promotions if len(p['dist_names'])>1})})")


if __name__ == "__main__":
    main()
